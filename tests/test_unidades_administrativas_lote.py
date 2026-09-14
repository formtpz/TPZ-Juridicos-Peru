import unittest
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from modulos.depuracion_comun import find_crc_column, normalize_crc, prepare_delivery_keys
from modulos.procesamiento_ua_lote import (
    OUTPUT_COLUMNS,
    SUMMARY_COUNT_COLUMN,
    build_excel,
    extract_crc_components,
    process_report,
)


class UnidadesAdministrativasLoteTests(unittest.TestCase):
    def setUp(self):
        self.crc_a = "150133490380200101010017"
        self.crc_b = "150133490380200101020018"
        self.crc_other_ubigeo = "150142490380200101010017"
        self.crc_common = "150133490380200101019997"

    def report(self):
        return pd.DataFrame(
            {
                "Departamento": ["Lima"] * 4,
                "Provincia": ["Lima"] * 4,
                "Distrito": ["SJM", "SJM", "VES", "SJM"],
                "Ubigeo": [150133, 150133, 150142, 150133],
                "Código de Referencia Catastral": [
                    self.crc_a,
                    self.crc_b,
                    self.crc_other_ubigeo,
                    self.crc_common,
                ],
                "Número de Partida Registral": ["P1", "P2", "P3", "P4"],
            }
        )

    def delivery_keys(self):
        deliveries = pd.DataFrame(
            {
                "poligono": ["SJM-15", "SJM-15", "SJM-15"],
                "ubigeo": [150133, 150133, 150133],
                "concat_sec": [49038, 49038, 49038],
                "fuente_entrega": ["campo", "campo", "cofopri"],
            }
        )
        return prepare_delivery_keys(deliveries, ["SJM-15"])

    def test_normalize_crc_text_number_decimal_and_scientific(self):
        self.assertEqual(normalize_crc(f"  {self.crc_a}  "), self.crc_a)
        self.assertEqual(normalize_crc(self.crc_a + ".0"), self.crc_a)
        scientific = "1.50133490380200101010017E+23"
        self.assertEqual(normalize_crc(scientific), self.crc_a)
        self.assertEqual(normalize_crc(123), "000000000000000000000123")
        self.assertIsNone(normalize_crc("ABC"))

    def test_extract_components(self):
        parts = extract_crc_components(self.crc_a)
        self.assertEqual(parts["Sector"], "49")
        self.assertEqual(parts["Manzana"], "038")
        self.assertEqual(parts["Lote"], "020")
        self.assertEqual(parts["Unidad"], "001")
        self.assertEqual(parts["Concat"], "49038020")

    def test_crc_column_fallback_only_accepts_generic_headers(self):
        generic = pd.DataFrame(columns=["Field1", "Field2", "Field3", "Field4", "Field5"])
        self.assertEqual(find_crc_column(generic), "Field5")
        with self.assertRaises(ValueError):
            find_crc_column(pd.DataFrame(columns=["A", "B", "C", "D", "E"]))

    def test_filter_count_order_exclusion_and_no_duplicate_rows(self):
        filtered, summary = process_report(self.report(), self.delivery_keys())
        self.assertEqual(len(filtered), 2)
        self.assertEqual(filtered.columns.tolist(), OUTPUT_COLUMNS)
        self.assertEqual(filtered["UC"].tolist(), [2, 2])
        self.assertNotIn(self.crc_common, filtered["Código de Referencia Catastral"].tolist())
        self.assertNotIn(self.crc_other_ubigeo, filtered["Código de Referencia Catastral"].tolist())
        self.assertEqual(summary.iloc[0].to_dict(), {"Row Labels": "49038020", SUMMARY_COUNT_COLUMN: 2})
        self.assertEqual(summary.iloc[-2]["Row Labels"], "(blank)")
        self.assertEqual(summary.iloc[-1].to_dict(), {"Row Labels": "Grand Total", SUMMARY_COUNT_COLUMN: 2})

    def test_workbook_structure(self):
        filtered, summary = process_report(self.report(), self.delivery_keys())
        workbook = load_workbook(BytesIO(build_excel(filtered, summary)), read_only=False)
        self.assertEqual(workbook.sheetnames, ["Filtrado", "Sheet1"])
        self.assertEqual([cell.value for cell in workbook["Filtrado"][1]], OUTPUT_COLUMNS)
        self.assertEqual(workbook["Filtrado"].auto_filter.ref, "A1:O3")


if __name__ == "__main__":
    unittest.main()
